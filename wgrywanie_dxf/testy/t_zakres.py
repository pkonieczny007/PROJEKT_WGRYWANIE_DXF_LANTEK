# -*- coding: utf-8 -*-
"""Zakres wgrywania: referencja + ZAKUPY=blacha (wl/wyl) + warunki 'kolumna = wartosc'
dodawane plusem + pomijanie MAGAZYN (tylko kolumny na lewo od referencji).
Tylko odczyt wykazu (i SELECT z bazy przy kontroli), bez wgrywania.

    python testy/t_zakres.py [NUMER]
"""
import glob
import io
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

TMP = Path(tempfile.gettempdir()) / "test_wgrywanie_dxf_zakres"
shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir()
os.environ["DATA_DIR"] = str(TMP / "data")
os.environ["LANTEK_DXF_DIR"] = str(TMP / "DXF")
os.environ["REJESTR_ZRODEL_DIR"] = str(TMP / "rejestr_zrodel")   # NIE prawdziwy na QNAP

from openpyxl import load_workbook  # noqa: E402

from app import app  # noqa: E402

NUMER = sys.argv[1] if len(sys.argv) > 1 else "75_8981"
WYKAZ = glob.glob("//qnap-energo/Technologia/2026/%s/ROBOCZY_*.xls*" % NUMER)[0]
c = app.test_client()
OK_ALL = True


def sprawdz(opis, warunek):
    global OK_ALL
    OK_ALL &= bool(warunek)
    print("  [%s] %s" % ("OK " if warunek else "ZLE", opis))


def stan():
    with c.session_transaction() as s:
        return json.loads((Path(s["folder"]) / "stan.json").read_text(encoding="utf-8"))


def kontrola(**pola):
    dane = {"akcja": "kontrola", "numer": NUMER, "wykaz": WYKAZ, "badaj_dxf": "",
            "folder_dxf": [str(TMP)], "ordref": ""}
    dane.update(pola)
    c.post("/dxf/", data=dane)
    c.post("/dxf/przebieg", data={"akcja": "kontrola"})
    return stan()


print("wykaz:", WYKAZ)
j = c.get("/dxf/api/wykaz_kolumny?wykaz=" + WYKAZ).get_json()
spr = next(k for k in j["naglowki"] if k["kol"] == j["domyslna"])
print("1. podpowiedzi kolumny '%s':" % spr["naglowek"], [(w["v"], w["n"]) for w in spr["wartosci"]])
sprawdz("domyslna kolumna warunku = sprawdzenie, wartosci z TEGO wykazu",
        spr["naglowek"].lower() == "sprawdzenie" and spr["wartosci"])

st = kontrola()
b = st["kontrola"]["bilans"]
print("\n2. bez pol zakresu (stary klient) -> domyslny:", st["kontrola"]["zakres_opis"], b)
sprawdz("domyslnie: referencja + blacha + bez MAGAZYN", "ZAKUPY = blacha" in st["kontrola"]["zakres_opis"]
        and "bez MAGAZYN" in st["kontrola"]["zakres_opis"])
pom = st["kontrola"]["pominiete"]
sprawdz("MAGAZYN pominiete (%d), zadne nie weszlo do pozycji" % len(pom),
        pom and not ({x["ref"] for x in pom} & {p["ref"] for p in st["podglad"]}))
wszystkie = b["wierszy"]

kol = {"kol": spr["kol"], "naglowek": spr["naglowek"]}
w1 = spr["wartosci"][0]["v"]
st = kontrola(zakupy_blacha="on", pomin_magazyn="on",
              zakres_json=json.dumps([dict(kol, wartosc=w1.upper())]))
n1 = st["kontrola"]["bilans"]["wierszy"]
print("\n3. sprawdzenie = %s:" % w1.upper(), st["kontrola"]["zakres_opis"], "->", n1)
sprawdz("warunek zaweza zakres (bez wielkosci liter): %d < %d" % (n1, wszystkie),
        0 < n1 < wszystkie and n1 == spr["wartosci"][0]["n"])
w2 = spr["wartosci"][1]["v"]
st = kontrola(zakupy_blacha="on", pomin_magazyn="on",
              zakres_json=json.dumps([dict(kol, wartosc=w1), dict(kol, wartosc=w2)]))
n2 = st["kontrola"]["bilans"]["wierszy"]
print("4. + sprawdzenie = %s:" % w2, st["kontrola"]["zakres_opis"], "->", n2)
sprawdz("drugi warunek na tej samej kolumnie poszerza zakres: %d > %d" % (n2, n1), n2 > n1)

st = kontrola(pomin_magazyn="on", zakres_json="[]")          # blacha WYLACZONA
n3 = st["kontrola"]["bilans"]["wierszy"]
print("5. bez filtra ZAKUPY:", st["kontrola"]["zakres_opis"], "->", n3)
sprawdz("bez ZAKUPY=blacha wiecej wierszy (%d >= %d)" % (n3, wszystkie), n3 >= wszystkie)

st = kontrola(zakupy_blacha="on", zakres_json="[]")          # MAGAZYN NIE pomijany
sprawdz("bez pomijania MAGAZYN: %d wierszy, 0 pominietych" % st["kontrola"]["bilans"]["wierszy"],
        st["kontrola"]["bilans"]["magazyn"] == 0 and st["kontrola"]["bilans"]["wierszy"] == wszystkie + len(pom))

st = kontrola()
x = c.get("/dxf/eksport")
wb = load_workbook(io.BytesIO(x.data))
sprawdz("Excel: zakladka 'Pominiete MAGAZYN' (%s)" % wb.sheetnames, "Pominiete MAGAZYN" in wb.sheetnames)
h = c.get("/dxf/przebieg").get_data(as_text=True)
sprawdz("strona kontroli pokazuje zakres i pominiete", "Zakres:" in h and "pominiete MAGAZYN" in h)

print("\n6. 'MAGAZYN 2szt' = czesc z magazynu: wgrywamy, ale oznaczone")
from services import wykaz as W  # noqa: E402
f2 = glob.glob("//qnap-energo/Technologia/2026/57_5406/*oboczy_*.xls*")
if f2:
    w = W.wczytaj(f2[0])
    cz = [p for p in w.pozycje if p.magazyn]
    print("     57_5406:", [(p.wiersz, p.magazyn, p.ilosc) for p in cz[:3]])
    sprawdz("czesciowe zostaja w zakresie (%d), pelne 'MAGAZYN' pominiete (%d)"
            % (len(cz), len(w.pominiete)), cz and not ({p.ref for p in cz} & {x["ref"] for x in w.pominiete}))
    st = kontrola(wykaz=f2[0])
    h = c.get("/dxf/przebieg").get_data(as_text=True)
    sprawdz("strona: wiersze podswietlone (%d) + licznik w zakresie" % h.count('magazyn" data-poziom'),
            h.count('magazyn" data-poziom') == len(cz) and "czesciowo z magazynu" in h)
    sprawdz("uwaga MAGAZYN_CZESCIOWO przy pozycjach",
            st["kontrola"]["kody"].get("MAGAZYN_CZESCIOWO") == len(cz))
else:
    print("     (brak wykazu 57_5406 - pomijam)")

print("\nWYNIK:", "OK" if OK_ALL else "SA BLEDY")
sys.exit(0 if OK_ALL else 1)
