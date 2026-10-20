# -*- coding: utf-8 -*-
"""Rejestr zrodel DXF (Excel): dopisywanie, backup, plik zajety -> kolejka.
W katalogu tymczasowym - prawdziwego rejestru na QNAP nie dotyka.

    python testy/t_rejestr_zrodel.py
"""
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from services import rejestr_zrodel as R  # noqa: E402

KAT = Path(tempfile.gettempdir()) / "test_rejestr_zrodel"
shutil.rmtree(KAT, ignore_errors=True)
OK_ALL = True


def sprawdz(opis, warunek):
    global OK_ALL
    OK_ALL &= bool(warunek)
    print("  [%s] %s" % ("OK " if warunek else "ZLE", opis))


st = {"numer": "82_1129", "wykaz": "//q/Technologia/2026/82_1129/ROBOCZY_x.xlsx",
      "ordref": "82_1129_BLACHY", "folder_przebiegu": "C:/data/82_1129_x",
      "do_wgrania": [{"dxf_zrodlo": "//q/T/82/DXF_A/a1.dxf"}, {"dxf_zrodlo": "//q/T/82/DXF_A/a2.dxf"},
                     {"dxf_zrodlo": "//q/T/82/DXF_B/b1.dxf"}, {"dxf_zrodlo": None}]}
w = R.wiersze_wgrania(st, "polecen 9, BLEDOW 0", "ponowne wgranie")
sprawdz("jeden wiersz na folder: %s" % [(x[0][-5:], x[6]) for x in w],
        len(w) == 2 and [x[6] for x in w] == [2, 1])
sprawdz("kolejnosc kolumn: folder, wykaz, zamowienie, uwagi",
        w[0][1] == st["wykaz"] and w[0][2] == "82_1129_BLACHY" and w[0][3] == "ponowne wgranie")

r = R.dopisz(KAT, w)
sprawdz("pierwszy zapis: plik + folder backup", r["ok"] and (KAT / R.PLIK).is_file() and (KAT / "backup").is_dir())
r = R.dopisz(KAT, w[:1])
sprawdz("drugi zapis: stara wersja w backup (%s)" % Path(r["backup"] or "-").name,
        r["ok"] and r["backup"] and Path(r["backup"]).is_file())
ws = load_workbook(str(KAT / R.PLIK)).worksheets[0]
sprawdz("naglowek + 3 wiersze (%d)" % (ws.max_row - 1),
        ws.cell(1, 1).value == R.KOLUMNY[0] and ws.max_row == 4)

# plik zajety (otwarty w Excelu) -> kolejka, potem dopisanie
oryg = R.load_workbook


def zajety(*a, **k):
    raise PermissionError("[Errno 13] Permission denied")


R.load_workbook = zajety
r = R.dopisz(KAT, w[:1])
R.load_workbook = oryg
sprawdz("plik zajety: nie zapisany, w kolejce %d" % r["oczekuje"],
        not r["ok"] and r["oczekuje"] == 1 and (KAT / R.OCZEKUJACE).is_file())
r = R.dopisz(KAT, w[1:])
ws = load_workbook(str(KAT / R.PLIK)).worksheets[0]
sprawdz("nastepny zapis dopisal kolejke + nowy wiersz (%d wierszy)" % (ws.max_row - 1),
        r["ok"] and ws.max_row == 6 and not (KAT / R.OCZEKUJACE).exists())

print("\nWYNIK:", "OK" if OK_ALL else "SA BLEDY")
sys.exit(0 if OK_ALL else 1)
