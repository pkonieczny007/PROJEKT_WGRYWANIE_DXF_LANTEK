# -*- coding: utf-8 -*-
"""Eksport podgladu kontroli do Excela: zakladka 'Pozycje' i 'Materialy'.

Zrodlo = stan.json przebiegu (to samo, co widac na stronie po ostatniej
kontroli). Zmiany reczne podswietlone na fioletowo, jak na stronie.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KOLOR_POZIOMU = {"OK": "E3F2E8", "UWAGA": "FFF1DC", "STOP": "FBE0DE"}
RECZNA = PatternFill("solid", fgColor="E9DDFD")
NAGLOWEK = PatternFill("solid", fgColor="DDE3EA")
POGRUB = Font(bold=True)

# (naglowek, klucz w podgladzie albo funkcja, pole 'reczne' do podswietlenia)
KOLUMNY = [
    ("poziom", "poziom", None),
    ("wiersz wykazu", "wiersz", None),
    ("strona L/P", "strona", None),
    ("referencja (do LANTEKA)", "ref", "ref"),
    ("plik DXF", "ref_dxf", None),
    ("referencja w wykazie", "ref_wykaz", None),
    ("gatunek", "gatunek", None),
    ("grubosc_gatunek", lambda p: "%s_%s" % (
        ("%g" % p["grubosc"]) if isinstance(p.get("grubosc"), (int, float)) else (p.get("grubosc") or ""),
        p.get("gatunek") or ""), None),
    ("MatRef", "material", "material"),
    ("grubosc [mm]", "grubosc", "grubosc"),
    ("ilosc", "ilosc", "ilosc"),
    ("kod technologii", "kod", None),
    ("maszyna", "maszyna", "maszyna"),
    ("skad maszyna", "maszyna_zrodlo", None),
    ("LANTEK sam - maszyna", "maszyna_lantek", None),
    ("LANTEK sam - ruta", "ruta", None),
    ("ruta reczna", "ruta_reczna", "ruta"),
    ("marszruta (przewidziana)", "marszruta", None),
    ("poprawka marszruty", lambda p: "tak" if p.get("popraw_marszrute") else "", None),
    ("stan DXF", "stan_dxf", None),
    ("sciezka DXF", "dxf_zrodlo", None),
    ("juz w bazie", lambda p: "tak" if p.get("w_bazie") else "", None),
    ("w zamowieniach", lambda p: ", ".join(p.get("w_zamowieniach") or []), None),
    ("zmiany reczne (bylo -> jest)", lambda p: "; ".join(
        "%s: %s -> %s" % (k, v[0], v[1]) for k, v in (p.get("reczne") or {}).items()), None),
    ("czesciowo z magazynu", "magazyn", None),
    ("uwagi", lambda p: "\n".join(p.get("uwagi") or []), None),
]
MAGAZYN = PatternFill("solid", fgColor="FFE7A8")


def _zmienione(p) -> set:
    """Ktore pola tej pozycji czlowiek zmienil recznie."""
    out = set((p.get("reczne") or {}).keys())
    if p.get("ref") and p.get("ref_dxf") and p["ref"] != p["ref_dxf"]:
        out.add("ref")
    if p.get("maszyna_zrodlo") in ("pozycja", "ruta"):
        out.add("maszyna")
    if p.get("ruta_reczna"):
        out.add("ruta")
    return out


def _szerokosci(ws, maks=60):
    for i, kol in enumerate(ws.columns, start=1):
        dl = max((len(str(c.value).split("\n")[0]) for c in kol if c.value is not None),
                 default=6)
        ws.column_dimensions[get_column_letter(i)].width = min(max(dl + 2, 8), maks)


def _naglowek(ws, nazwy):
    ws.append(nazwy)
    for c in ws[1]:
        c.font, c.fill = POGRUB, NAGLOWEK
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def zbuduj(st) -> Workbook:
    k = st.get("kontrola") or {}
    wb = Workbook()

    ws = wb.active
    ws.title = "Pozycje"
    _naglowek(ws, [n for n, _k, _r in KOLUMNY])
    for p in st.get("podglad") or []:
        zm = _zmienione(p)
        ws.append([(kl(p) if callable(kl) else p.get(kl)) for _n, kl, _r in KOLUMNY])
        r = ws.max_row
        kolor = KOLOR_POZIOMU.get(p.get("poziom"))
        if kolor:
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=kolor)
        if p.get("magazyn"):                 # 'MAGAZYN 2szt' - caly wiersz na bursztynowo
            for cell in ws[r]:
                cell.fill = MAGAZYN
        for i, (_n, _kl, pole) in enumerate(KOLUMNY, start=1):
            if pole and pole in zm:
                ws.cell(r, i).fill = RECZNA
                ws.cell(r, i).font = POGRUB
    ws.auto_filter.ref = ws.dimensions
    _szerokosci(ws)

    wm = wb.create_sheet("Materialy")
    _naglowek(wm, ["grubosc [mm]", "gatunek", "MatRef w LANTEKU", "pozycji", "sztuk",
                   "jest w kartotece", "maszyna", "maszyna domyslna", "wyjatkow przy pozycjach"])
    poz = szt = 0
    for m in k.get("materialy") or []:
        wkart = m.get("w_kartotece")
        wm.append([m.get("grubosc"), m.get("gatunek"), m.get("material"), m.get("pozycji"),
                   m.get("sztuk"), "tak" if wkart else ("NIE MA" if wkart is False else "?"),
                   m.get("maszyna"), m.get("maszyna_domyslna"), m.get("wyjatkow") or 0])
        if m.get("maszyna") and m.get("maszyna") != m.get("maszyna_domyslna"):
            wm.cell(wm.max_row, 7).fill = RECZNA
            wm.cell(wm.max_row, 7).font = POGRUB
        poz += m.get("pozycji") or 0
        szt += m.get("sztuk") or 0
    wm.append(["SUMA", "", "%d grup" % len(k.get("materialy") or []), poz, szt])
    for c in wm[wm.max_row]:
        c.font = POGRUB
    b = k.get("bilans") or {}
    if b:
        wm.append([])
        wm.append(["bilans", "wierszy blachy %s + LP %s = oczekiwanych %s; pozycji po kontroli %s"
                   % (b.get("wierszy"), b.get("lp"), b.get("oczekiwanych"), b.get("pozycji"))])
    if k.get("zakres_opis"):
        wm.append(["zakres", k["zakres_opis"]])
    _szerokosci(wm)

    pom = k.get("pominiete") or []
    if pom:
        wp = wb.create_sheet("Pominiete MAGAZYN")
        _naglowek(wp, ["wiersz wykazu", "referencja", "dlaczego"])
        for x in pom:
            wp.append([x.get("wiersz"), x.get("ref"), x.get("powod")])
        _szerokosci(wp)
    return wb


def do_pamieci(st) -> io.BytesIO:
    buf = io.BytesIO()
    zbuduj(st).save(buf)
    buf.seek(0)
    return buf


def nazwa_pliku(st) -> str:
    return "PODGLAD_%s.xlsx" % (st.get("numer") or "bez_numeru").replace(" ", "_")
