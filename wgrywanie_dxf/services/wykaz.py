# -*- coding: utf-8 -*-
"""Czytanie wykazu roboczego (pierwsza zakladka) -> lista pozycji BLACHA.

Zasady (z /wykaz-roboczy i EKSTRAKTOR_DXF V3):
  * kolumny szukamy PO NAGLOWKU, nie po literze; przy duplikacie wygrywa
    kolumna NAJBARDZIEJ Z LEWEJ (rozbicie 'Bezeichnung' kopiuje naglowek w prawo),
  * wiersz naglowka = wiersz z najwieksza liczba wypelnionych komorek (skan 15),
  * zakres = ZAKUPY zawiera 'blach' (case-insensitive),
  * NAZWA/ZAKUPY bywaja FORMULAMI - czytamy data_only=True (cache Excela);
    gdy cache pusty, mowimy o tym GLOSNO zamiast wgrywac puste referencje.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .wspolne import do_liczby, norm_naglowek, rozbierz_referencje, sciezka

MAPOWANIE = json.loads((Path(__file__).resolve().parent.parent / "mapowanie.json")
                       .read_text(encoding="utf-8"))
ALIASY = MAPOWANIE["kolumny"]


@dataclass
class Pozycja:
    wiersz: int
    ref: str = ""                 # referencja = nazwa produktu w LANTEKU
    zeinr: str = ""
    posn: str = ""
    ilosc: float | None = None    # z kolumny ilosci
    ilosc_ref: int | None = None  # z czlonu 'Nst' referencji
    grubosc_kol: float | None = None   # kolumna grubosc_poprawiona (bywa jej brak)
    gatunek_kol: str = ""              # kolumna gatunek_poprawiony (bywa jej brak)
    grubosc_ref: float | None = None   # 1. czlon referencji
    gatunek_ref: str = ""              # 2. czlon referencji
    zrodlo: str = "referencja"         # ktore zrodlo obowiazuje: referencja | kolumny
    kod_kolumna: str = ""         # kolumna 'Giecie'
    kod_ref: str | None = None    # technologia odczytana z referencji
    sprawdzenie: str = ""
    propozycja: str = ""
    indeks: str = ""
    bezeichnung: str = ""
    uwagi: list = field(default_factory=list)
    # ref_wykaz - jak stoi w wykazie; ref_dxf - nazwa pliku DXF (po rozbiciu LP na
    # L i P, ale PRZED poprawka czlowieka); ref - co idzie do LANTEKA
    ref_wykaz: str = ""
    ref_dxf: str = ""
    strona: str = ""              # 'L' / 'P' gdy pozycja powstala z wiersza LP
    reczne: dict = field(default_factory=dict)   # {'grubosc': [bylo, jest], ...} zmiany czlowieka
    magazyn: str = ""             # 'MAGAZYN 2szt' - czesc z magazynu, ilosc pomniejszona

    @property
    def kod(self) -> str:
        """Do marszruty bierzemy kod Z REFERENCJI - to ona idzie do LANTEKA."""
        zrodlo = self.kod_ref if self.kod_ref is not None else self.kod_kolumna
        return (zrodlo or "").strip().upper()

    @property
    def grubosc(self):
        """Obowiazujaca grubosc. Domyslnie z referencji; gdy jej brak - z kolumny."""
        if self.zrodlo == "kolumny":
            return self.grubosc_kol if self.grubosc_kol is not None else self.grubosc_ref
        return self.grubosc_ref if self.grubosc_ref is not None else self.grubosc_kol

    @property
    def gatunek(self) -> str:
        """Obowiazujacy gatunek. Domyslnie z referencji; gdy jej brak - z kolumny."""
        if self.zrodlo == "kolumny":
            return self.gatunek_kol or self.gatunek_ref or ""
        return self.gatunek_ref or self.gatunek_kol or ""

    @property
    def rozjazd_grubosci(self) -> bool:
        return (self.grubosc_kol is not None and self.grubosc_ref is not None
                and abs(float(self.grubosc_kol) - float(self.grubosc_ref)) > 0.001)

    @property
    def rozjazd_gatunku(self) -> bool:
        a = (self.gatunek_kol or "").replace(" ", "").upper()
        b = (self.gatunek_ref or "").replace(" ", "").upper()
        return bool(a and b and a != b)


@dataclass
class Wykaz:
    plik: str
    zakladka: str
    wiersz_naglowka: int
    naglowki: dict                 # indeks kolumny -> tekst
    kolumny: dict                  # rola -> indeks
    kandydaci_nazwy: dict          # naglowek -> indeks
    pozycje: list                  # Pozycja (tylko BLACHA)
    wierszy: int
    blacha: int
    kolumna_nazwy: str = ""
    kolumna_ilosci: str = ""
    formuly_puste: int = 0         # ile referencji puste mimo formuly
    # zakres: co zostalo poza nim (zeby czlowiek widzial, czego NIE wgrywamy)
    pominiete: list = field(default_factory=list)   # [{'wiersz','ref','powod'}] - MAGAZYN
    bez_referencji: int = 0        # wierszy w zakresie bez referencji (nie wgrywamy)
    poza_zakresem: int = 0         # wierszy z referencja, ktore nie spelnily warunkow
    zakres_opis: str = ""
    zakres_uwagi: list = field(default_factory=list)


ZAKRES = MAPOWANIE.get("zakres", {})
_MAGAZYN_KOL = {norm_naglowek(k) for k in ZAKRES.get("kolumny_magazyn", [])}
_MAGAZYN_SLOWO = (ZAKRES.get("slowo_magazyn") or "magazyn").lower()


def zakres_domyslny() -> dict:
    """Referencja + ZAKUPY=blacha, bez dodatkowych warunkow, MAGAZYN pomijany."""
    return {"zakupy_blacha": True, "warunki": [], "pomin_magazyn": True}


def _rowne(a, b) -> bool:
    """Porownanie wartosci komorek bez wielkosci liter i spacji na brzegach
    ('Z BAZY' == 'z bazy')."""
    return _txt(a).lower() == _txt(b).lower()


def _warunki_po_kolumnach(warunki, naglowki) -> tuple:
    """[{'kol': idx, 'naglowek': txt, 'wartosc': v}] -> ({idx: [wartosci]}, uwagi).

    Indeks kolumny sprawdzamy z naglowkiem - gdyby wykaz sie zmienil, szukamy
    kolumny po tekscie naglowka. Ta sama kolumna kilka razy = LUB (zakres rosnie),
    rozne kolumny = I."""
    grupy, uwagi = {}, []
    for w in warunki or []:
        nag = _txt(w.get("naglowek"))
        try:
            idx = int(w.get("kol"))
        except (TypeError, ValueError):
            idx = None
        if idx is None or naglowki.get(idx) != nag:
            idx = next((i for i, h in naglowki.items() if h == nag), None)
        if idx is None:
            uwagi.append("warunek pominiety - w wykazie nie ma kolumny '%s'" % nag)
            continue
        grupy.setdefault(idx, []).append(_txt(w.get("wartosc")))
    return grupy, uwagi


def opis_zakresu(zakres, naglowki=None) -> str:
    zakres = zakres or zakres_domyslny()
    cz = ["referencja wpisana"]
    if zakres.get("zakupy_blacha", True):
        cz.append("ZAKUPY = blacha")
    po_kol = {}
    for w in zakres.get("warunki") or []:
        po_kol.setdefault(_txt(w.get("naglowek")), []).append(_txt(w.get("wartosc")) or "(puste)")
    for k, v in po_kol.items():
        cz.append("%s = %s" % (k, " LUB ".join(v)))
    if zakres.get("pomin_magazyn", True):
        cz.append("bez MAGAZYN")
    return " + ".join(cz)


def _mapa_naglowkow(hdr):
    poz = {}
    for i, h in enumerate(hdr):
        k = norm_naglowek(h)
        if k:
            poz.setdefault(k, i)           # LEWA wygrywa
    return poz


def _znajdz(poz, aliasy):
    """Pierwszy alias z listy, ktory ma trafienie; w razie kilku - lewa kolumna."""
    for a in aliasy:
        ka = norm_naglowek(a)
        traf = [i for k, i in poz.items() if k == ka]
        if not traf:
            traf = [i for k, i in poz.items() if k.startswith(ka)]
        if traf:
            return min(traf)
    return None


def _wiersz_naglowka(ws, limit=15):
    najlepszy, ile_najlepszy = 1, -1
    for nr, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), start=1):
        ile = sum(1 for c in row if c not in (None, ""))
        if ile > ile_najlepszy:
            najlepszy, ile_najlepszy = nr, ile
    return najlepszy


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def wczytaj(plik, kolumna_nazwy: str | None = None, kolumna_ilosci: str | None = None,
            zrodlo: str = "referencja", zakres: dict | None = None) -> Wykaz:
    """kolumna_nazwy / kolumna_ilosci: tekst naglowka wskazany przez uzytkownika
    (None = wybor automatyczny wg mapowania).

    zrodlo: skad brac grubosc i gatunek - 'referencja' (domyslnie) albo 'kolumny'.
    Referencja niesie oba czlony zawsze i to ona idzie do LANTEKA; kolumny
    kontrolne istnieja tylko w wykazach zrobionych przez /wykaz-roboczy.

    zakres: {'zakupy_blacha': bool, 'warunki': [{'kol','naglowek','wartosc'}],
    'pomin_magazyn': bool} - patrz zakres_domyslny(). Wiersz bez referencji nigdy
    nie idzie do LANTEKA (liczymy go tylko, zeby wylapac nieprzeliczone formuly).
    """
    zakres = zakres or zakres_domyslny()
    p = sciezka(plik)
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    nr_hdr = _wiersz_naglowka(ws)
    rows = ws.iter_rows(min_row=nr_hdr, values_only=True)
    hdr = next(rows)
    poz = _mapa_naglowkow(hdr)
    naglowki = {i: _txt(h) for i, h in enumerate(hdr) if h is not None}

    kol = {rola: _znajdz(poz, al) for rola, al in ALIASY.items()}

    # kandydaci na kolumne referencji: kazda kolumna, ktorej naglowek zaczyna sie od 'nazwa'
    kandydaci = {naglowki[i]: i for k, i in poz.items()
                 if k.startswith("nazwa") and i in naglowki}
    # kolumna wskazana przez czlowieka (lista z naglowkow albo wpisana): dowolny
    # naglowek, bez wielkosci liter i kropek ('Anz.' = 'ANZ'), wygrywa lewa;
    # gdy takiej nie ma - zostaje wybor automatyczny i mowimy o tym
    kol_uwagi = []
    for rola, wskazana in (("nazwa", kolumna_nazwy), ("ilosc", kolumna_ilosci)):
        if not wskazana:
            continue
        traf = sorted(i for i, h in naglowki.items() if norm_naglowek(h) == norm_naglowek(wskazana))
        if traf:
            kol[rola] = traf[0]
        else:
            kol_uwagi.append("w wykazie nie ma kolumny '%s' - %s wybrana automatycznie (%s)"
                             % (wskazana, "referencja" if rola == "nazwa" else "ilosc",
                                naglowki.get(kol.get(rola), "brak")))

    zak_blacha = zakres.get("zakupy_blacha", True)
    brak = [r for r in (("zakupy", "nazwa") if zak_blacha else ("nazwa",))
            if kol.get(r) is None]
    if brak:
        wb.close()
        raise ValueError("w wykazie brak kolumn: %s (pierwsza zakladka '%s', "
                         "wiersz naglowka %d)" % (", ".join(brak), ws.title, nr_hdr))
    i_zak, i_naz = kol.get("zakupy"), kol["nazwa"]
    warunki, zakres_uwagi = _warunki_po_kolumnach(zakres.get("warunki"), naglowki)
    # MAGAZYN: kolumny o tych naglowkach, ale TYLKO na lewo od kolumny z referencja -
    # na prawo bywaja kopie/duplikaty kolumn (decyzja uzytkownika)
    i_mag = [i for i, h in naglowki.items()
             if i < i_naz and norm_naglowek(h) in _MAGAZYN_KOL]

    def g(r, rola):
        i = kol.get(rola)
        return r[i] if i is not None and i < len(r) else None

    def kom(r, i):
        return r[i] if i is not None and i < len(r) else None

    pozycje, n = [], 0
    pominiete, bez_ref, poza = [], set(), 0
    for nr, r in enumerate(rows, start=nr_hdr + 1):
        n += 1
        if zak_blacha and "blach" not in _txt(kom(r, i_zak)).lower():
            continue
        ref = _txt(kom(r, i_naz))
        if warunki and not all(any(_rowne(kom(r, i), v) for v in wart)
                               for i, wart in warunki.items()):
            if ref:
                poza += 1
            continue
        if not ref:
            bez_ref.add(nr)
            continue
        # MAGAZYN: sam napis = caly element z magazynu -> nie wgrywamy;
        # z liczba ('MAGAZYN 2szt', '2 SZT', '2 sztuki') = czesc z magazynu,
        # ilosc pomniejszona -> wgrywamy, ale pozycja jest wyraznie oznaczona
        magazyn_txt = ""
        i_traf = next((i for i in i_mag
                       if _MAGAZYN_SLOWO in _txt(kom(r, i)).lower()), None)
        if i_traf is not None:
            tekst = _txt(kom(r, i_traf))
            if re.search(r"\d", tekst):
                magazyn_txt = "%s: %s" % (naglowki[i_traf], tekst[:80])
            elif zakres.get("pomin_magazyn", True):
                pominiete.append({"wiersz": nr, "ref": ref, "powod": "%s: %s"
                                  % (naglowki[i_traf], tekst[:80])})
                continue
        rozb = rozbierz_referencje(ref)
        pozycje.append(Pozycja(
            wiersz=nr, ref=ref,
            zeinr=_txt(g(r, "zeinr")) or (rozb["zeinr"] or ""),
            posn=_txt(g(r, "posn")) or (rozb["posn"] or ""),
            ilosc=do_liczby(g(r, "ilosc")),
            ilosc_ref=rozb["ilosc"],
            grubosc_kol=do_liczby(g(r, "grubosc")),
            gatunek_kol=_txt(g(r, "gatunek")),
            grubosc_ref=rozb["grubosc"],
            gatunek_ref=rozb["gatunek"] or "",
            zrodlo=zrodlo,
            kod_kolumna=_txt(g(r, "giecie")),
            kod_ref=rozb["technologia"],
            sprawdzenie=_txt(g(r, "sprawdzenie")),
            propozycja=_txt(g(r, "propozycja")),
            indeks=_txt(g(r, "index")),
            bezeichnung=_txt(g(r, "bezeichnung")),
            ref_wykaz=ref, ref_dxf=ref, magazyn=magazyn_txt,
        ))
    wb.close()

    w = Wykaz(plik=str(p), zakladka=ws.title, wiersz_naglowka=nr_hdr, naglowki=naglowki,
              kolumny=kol, kandydaci_nazwy=kandydaci, pozycje=pozycje,
              wierszy=n, blacha=len(pozycje),
              kolumna_nazwy=naglowki.get(i_naz, ""),
              kolumna_ilosci=naglowki.get(kol.get("ilosc"), ""),
              pominiete=pominiete, bez_referencji=len(bez_ref), poza_zakresem=poza,
              zakres_opis=opis_zakresu(zakres), zakres_uwagi=kol_uwagi + zakres_uwagi)

    # czy puste referencje biora sie z nieprzeliczonej formuly?
    if bez_ref:
        w.formuly_puste = _policz_formuly(p, nr_hdr, i_naz, bez_ref)
    return w


_cache_kolumn = {}


def kolumny_i_wartosci(plik, max_wartosci=80) -> dict:
    """Do UI zakresu: naglowki pierwszej zakladki i wartosci kazdej kolumny
    (bez wielkosci liter, z liczba wierszy). Kolumna z wiecej niz max_wartosci
    roznymi wartosciami dostaje pusta liste - wtedy wartosc wpisuje sie recznie.
    Wynik trzymamy w pamieci po (sciezka, data zmiany pliku)."""
    p = sciezka(plik)
    klucz = (str(p).lower(), p.stat().st_mtime)
    if klucz in _cache_kolumn:
        return _cache_kolumn[klucz]
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    nr_hdr = _wiersz_naglowka(ws)
    rows = ws.iter_rows(min_row=nr_hdr, values_only=True)
    hdr = next(rows)
    naglowki = {i: _txt(h) for i, h in enumerate(hdr) if _txt(h)}
    licz = {i: {} for i in naglowki}
    for r in rows:
        for i in naglowki:
            if licz[i] is None or i >= len(r):
                continue
            v = _txt(r[i])
            if not v:
                continue
            d = licz[i]
            k = v.lower()
            if k in d:
                d[k][1] += 1
            elif len(d) >= max_wartosci:
                licz[i] = None                       # za duzo roznych - bez podpowiedzi
            else:
                d[k] = [v, 1]
    wb.close()
    dom = norm_naglowek(ZAKRES.get("domyslna_kolumna_warunku", "sprawdzenie"))
    out = {"naglowki": [
        {"kol": i, "naglowek": h,
         "wartosci": [] if licz[i] is None else
         [{"v": v, "n": n} for v, n in sorted(licz[i].values(), key=lambda x: -x[1])]}
        for i, h in sorted(naglowki.items())],
        "domyslna": next((i for i, h in sorted(naglowki.items())
                          if norm_naglowek(h) == dom), None),
        "wiersz_naglowka": nr_hdr}
    _cache_kolumn.clear()                            # trzymamy tylko ostatni wykaz
    _cache_kolumn[klucz] = out
    return out


def _policz_formuly(p, nr_hdr, i_naz, wiersze) -> int:
    """openpyxl nie liczy formul; jesli w komorce siedzi '=' a cache jest pusty,
    to znaczy ze wykaz nie zostal przeliczony/zapisany w Excelu."""
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=False)
        ws = wb.worksheets[0]
        ile = 0
        for nr, r in enumerate(ws.iter_rows(min_row=nr_hdr + 1, values_only=True),
                               start=nr_hdr + 1):
            if nr in wiersze and i_naz < len(r) and isinstance(r[i_naz], str) \
                    and r[i_naz].startswith("="):
                ile += 1
        wb.close()
        return ile
    except Exception:
        return 0


def kandydaci_nazwy_wg_dxf(plik, nazwy_dxf) -> list:
    """Ile referencji z kazdej kolumny 'Nazwa*' trafia w pliki DXF - podpowiedz dla UI."""
    p = sciezka(plik)
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    nr_hdr = _wiersz_naglowka(ws)
    rows = ws.iter_rows(min_row=nr_hdr, values_only=True)
    hdr = next(rows)
    poz = _mapa_naglowkow(hdr)
    naglowki = {i: _txt(h) for i, h in enumerate(hdr) if h is not None}
    i_zak = _znajdz(poz, ALIASY["zakupy"])
    kand = {i: naglowki[i] for k, i in poz.items() if k.startswith("nazwa") and i in naglowki}
    licz = {i: 0 for i in kand}
    wypelnione = {i: 0 for i in kand}
    male = {str(n).lower() for n in nazwy_dxf}
    for r in rows:
        if i_zak is None or i_zak >= len(r):
            continue
        if "blach" not in _txt(r[i_zak]).lower():
            continue
        for i in kand:
            if i < len(r):
                v = _txt(r[i])
                if v:
                    wypelnione[i] += 1
                    if v.lower() in male:
                        licz[i] += 1
    wb.close()
    return sorted(({"kolumna": kand[i], "indeks": i, "trafien": licz[i],
                    "wypelnionych": wypelnione[i]} for i in kand),
                  key=lambda d: (-d["trafien"], -d["wypelnionych"]))
