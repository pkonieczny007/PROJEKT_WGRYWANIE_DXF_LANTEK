# -*- coding: utf-8 -*-
"""Rejestr zrodel wgranych DXF - wspolny plik Excel na QNAP.

Po kazdym wgraniu do LANTEKA dopisujemy po jednym wierszu na folder, z ktorego
faktycznie wzielismy pliki DXF. Dzieki temu kazdy wgrany element ma zrodlo.

Kolumny (kolejnosc ustalona z uzytkownikiem): folder DXF, wykaz, zamowienie
sprzedazy, uwagi - a za nimi szczegoly (data, zlecenie, ile plikow, wynik).

Plik bywa otwarty w Excelu u kogos innego. Wtedy wiersze odkladamy do
'oczekujace.jsonl' obok i dopisujemy przy nastepnym udanym zapisie - nic nie ginie.
Przed kazdym zapisem stara wersja idzie do backup/ (z data i godzina).
"""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .wspolne import sciezka

PLIK = "rejestr_zrodel_dxf.xlsx"
OCZEKUJACE = "oczekujace.jsonl"
KOLUMNY = ["folder DXF (zrodlo)", "wykaz", "zamowienie sprzedazy", "uwagi",
           "data wgrania", "zlecenie", "plikow DXF z folderu", "wynik importu",
           "folder przebiegu"]
SZEROKOSCI = [70, 70, 26, 50, 18, 12, 12, 40, 60]


def wiersze_wgrania(st, wynik_txt, uwagi) -> list:
    """Jeden wiersz na folder, z ktorego wzielismy pliki DXF do tego wgrania."""
    licz = {}
    for p in st.get("do_wgrania") or []:
        if p.get("dxf_zrodlo"):
            folder = str(Path(p["dxf_zrodlo"]).parent)
            licz[folder] = licz.get(folder, 0) + 1
    teraz = datetime.now().strftime("%Y-%m-%d %H:%M")
    return [[folder, st.get("wykaz") or "", st.get("ordref") or "", uwagi, teraz,
             st.get("numer") or "", ile, wynik_txt, st.get("folder_przebiegu") or ""]
            for folder, ile in sorted(licz.items())]


def _nowy() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejestr"
    ws.append(KOLUMNY)
    for i, c in enumerate(ws[1]):
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = SZEROKOSCI[i]
    ws.freeze_panes = "A2"
    return wb


def dopisz(katalog, wiersze) -> dict:
    """Dopisuje wiersze do rejestru. Zwraca {'ok', 'plik', 'backup', 'oczekuje', 'blad'}."""
    kat = sciezka(katalog)
    out = {"ok": False, "plik": str(kat / PLIK), "backup": None, "oczekuje": 0, "blad": ""}
    try:
        (kat / "backup").mkdir(parents=True, exist_ok=True)
    except OSError as e:
        out["blad"] = "brak dostepu do %s: %s" % (kat, e)
        return out
    plik, kolejka = kat / PLIK, kat / OCZEKUJACE

    zaleglosci = []
    if kolejka.is_file():
        for linia in kolejka.read_text(encoding="utf-8").splitlines():
            try:
                zaleglosci.append(json.loads(linia))
            except ValueError:
                pass
    wszystkie = zaleglosci + list(wiersze)
    if not wszystkie:
        out["ok"] = True
        return out

    try:
        if plik.is_file():
            kopia = kat / "backup" / ("%s_%s.xlsx" % (Path(PLIK).stem,
                                                      datetime.now().strftime("%Y%m%d_%H%M%S")))
            shutil.copy2(str(plik), str(kopia))
            out["backup"] = str(kopia)
            wb = load_workbook(str(plik))
        else:
            wb = _nowy()
        ws = wb.worksheets[0]
        for w in wszystkie:
            ws.append(w)
            ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = ws.dimensions
        wb.save(str(plik))
    except PermissionError as e:
        # plik otwarty w Excelu - odkladamy do kolejki, dopiszemy nastepnym razem
        with open(kolejka, "a", encoding="utf-8") as fh:
            for w in wiersze:
                fh.write(json.dumps(w, ensure_ascii=False) + "\n")
        out["oczekuje"] = len(zaleglosci) + len(wiersze)
        out["blad"] = "plik zajety (otwarty w Excelu?): %s" % e
        return out
    if kolejka.is_file():
        kolejka.unlink()
    out["ok"] = True
    out["dopisanych"] = len(wszystkie)
    return out
